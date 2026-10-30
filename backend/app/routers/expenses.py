from typing import List, Optional, Union
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from backend.app.database import get_db
from backend.app.models import serialize_doc, serialize_docs, build_id_query
from backend.app.schemas import ExpenseCreate, ExpenseUpdate, ExpenseOut
from backend.app.agent.tools import tool_add_expense, auto_categorize
from backend.app.services.anomaly import audit_and_generate_alerts

router = APIRouter(prefix="/api/expenses", tags=["Expenses"])

@router.get("", response_model=List[ExpenseOut])
async def list_expenses(
    month: Optional[str] = Query(None, description="Month in YYYY-MM format"),
    category: Optional[str] = Query(None),
    is_unusual: Optional[bool] = Query(None),
    search: Optional[str] = Query(None),
    limit: int = Query(100, le=500),
    user_id: str = "user_default",
    db = Depends(get_db)
):
    query: dict = {"user_id": user_id}
    if month:
        start_date = f"{month}-01"
        end_date = f"{month}-31"
        query["date"] = {"$gte": start_date, "$lte": end_date}
    if category and category.lower() != "all":
        query["category"] = category
    if is_unusual is not None:
        query["is_unusual"] = is_unusual
    if search:
        query["$or"] = [
            {"description": {"$regex": search, "$options": "i"}},
            {"category": {"$regex": search, "$options": "i"}}
        ]

    cursor = db.expenses.find(query).sort([("date", -1), ("_id", -1)]).limit(limit)
    docs = await cursor.to_list(length=limit)
    return serialize_docs(docs)

@router.post("", response_model=ExpenseOut)
async def create_expense(
    payload: ExpenseCreate,
    user_id: str = "user_default",
    db = Depends(get_db)
):
    res = await tool_add_expense(
        db=db,
        user_id=user_id,
        amount=payload.amount,
        category=payload.category,
        description=payload.description,
        date_str=payload.date,
        payment_method=payload.payment_method
    )
    query = {"$and": [{"user_id": user_id}, build_id_query(res["id"])]}
    obj = await db.expenses.find_one(query)
    return serialize_doc(obj)

@router.put("/{expense_id}", response_model=ExpenseOut)
async def update_expense(
    expense_id: str,
    payload: ExpenseUpdate,
    user_id: str = "user_default",
    db = Depends(get_db)
):
    query = {"$and": [{"user_id": user_id}, build_id_query(expense_id)]}
    exp = await db.expenses.find_one(query)
    if not exp:
        raise HTTPException(status_code=404, detail="Expense not found")

    update_fields = {}
    if payload.amount is not None:
        update_fields["amount"] = payload.amount
    if payload.category is not None:
        update_fields["category"] = payload.category
    if payload.description is not None:
        update_fields["description"] = payload.description
    if payload.date is not None:
        update_fields["date"] = payload.date
    if payload.payment_method is not None:
        update_fields["payment_method"] = payload.payment_method
    if payload.is_unusual is not None:
        update_fields["is_unusual"] = payload.is_unusual
    if payload.anomaly_reason is not None:
        update_fields["anomaly_reason"] = payload.anomaly_reason

    if update_fields:
        await db.expenses.update_one(query, {"$set": update_fields})
    
    updated_exp = await db.expenses.find_one(query)
    
    # Re-audit
    if updated_exp and "date" in updated_exp:
        await audit_and_generate_alerts(db, user_id, updated_exp["date"][:7])
        
    return serialize_doc(updated_exp)

@router.delete("/{expense_id}")
async def delete_expense(
    expense_id: str,
    user_id: str = "user_default",
    db = Depends(get_db)
):
    query = {"$and": [{"user_id": user_id}, build_id_query(expense_id)]}
    exp = await db.expenses.find_one(query)
    if not exp:
        raise HTTPException(status_code=404, detail="Expense not found")

    month_str = exp.get("date", "")[:7]
    await db.expenses.delete_one(query)

    # Re-audit
    if month_str:
        await audit_and_generate_alerts(db, user_id, month_str)
        
    return {"status": "success", "message": "Expense deleted"}

@router.post("/confirm-unusual/{expense_id}")
async def confirm_unusual(
    expense_id: str,
    user_id: str = "user_default",
    db = Depends(get_db)
):
    query = {"$and": [{"user_id": user_id}, build_id_query(expense_id)]}
    exp = await db.expenses.find_one(query)
    if not exp:
        raise HTTPException(status_code=404, detail="Expense not found")

    await db.expenses.update_one(
        query,
        {"$set": {"is_unusual": False, "anomaly_reason": "Verified & confirmed by user"}}
    )
    return {"status": "success", "message": "Expense confirmed"}


@router.post("/upload-csv")
async def upload_expenses_csv(
    file: UploadFile = File(...),
    user_id: str = "user_default",
    db = Depends(get_db)
):
    """
    Bulk-import expenses from Excel (.xlsx, .xls, .xlsm) or CSV spreadsheets.
    Supports multi-sheet selection, bank statements (SBI, HDFC, ICICI, etc.),
    flexible column names, Excel serial dates, timestamps, Indian currency (₹, Rs., /-),
    and automatic column type inference.
    """
    import csv
    import io
    import re
    from datetime import datetime, date, timedelta
    from dateutil import parser as date_parser

    filename = file.filename or "uploaded_spreadsheet.xlsx"
    raw_bytes = await file.read()
    if not raw_bytes:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    # ── 1. Read sheet/CSV into raw 2D matrix ─────────────────────────────────
    raw_table: list[list] = []
    lower_name = filename.lower()

    def parse_xlsx_bytes(b: bytes) -> list[list]:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(b), data_only=True)
        # Select worksheet with the most data
        best_sheet = None
        max_rows = -1
        for sheet in wb.worksheets:
            non_empty_cnt = 0
            for r in sheet.iter_rows(values_only=True):
                if any(c is not None and str(c).strip() != "" for c in r):
                    non_empty_cnt += 1
            if non_empty_cnt > max_rows:
                max_rows = non_empty_cnt
                best_sheet = sheet
        
        target_ws = best_sheet or wb.active
        table = []
        for r in target_ws.iter_rows(values_only=True):
            if any(c is not None and str(c).strip() != "" for c in r):
                table.append(list(r))
        return table

    def parse_xls_bytes(b: bytes) -> list[list]:
        import xlrd
        wb = xlrd.open_workbook(file_contents=b)
        best_sheet = None
        max_rows = -1
        for s_idx in range(wb.nsheets):
            sheet = wb.sheet_by_index(s_idx)
            if sheet.nrows > max_rows:
                max_rows = sheet.nrows
                best_sheet = sheet
        target_ws = best_sheet or wb.sheet_by_index(0)
        table = []
        for r_idx in range(target_ws.nrows):
            row = target_ws.row_values(r_idx)
            if any(c is not None and str(c).strip() != "" for c in row):
                table.append(row)
        return table

    if lower_name.endswith((".xlsx", ".xlsm", ".xltx")):
        try:
            raw_table = parse_xlsx_bytes(raw_bytes)
        except Exception as e_xlsx:
            try:
                raw_table = parse_xls_bytes(raw_bytes)
            except Exception:
                raise HTTPException(status_code=422, detail=f"Excel (.xlsx) read error: {e_xlsx}")
    elif lower_name.endswith(".xls"):
        try:
            raw_table = parse_xls_bytes(raw_bytes)
        except Exception as e_xls:
            try:
                raw_table = parse_xlsx_bytes(raw_bytes)
            except Exception:
                raise HTTPException(status_code=422, detail=f"Legacy Excel (.xls) read error: {e_xls}. Please save file as .xlsx or .csv.")
    else:
        # CSV or text file with multi-encoding and delimiter auto-detection
        decoded_text = None
        for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252", "iso-8859-1"):
            try:
                decoded_text = raw_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if not decoded_text:
            raise HTTPException(status_code=422, detail="Failed to decode CSV text. Please ensure UTF-8 encoding.")

        first_line = decoded_text.strip().splitlines()[0] if decoded_text.strip().splitlines() else ""
        delim = ","
        if ";" in first_line and first_line.count(";") > first_line.count(","):
            delim = ";"
        elif "\t" in first_line and first_line.count("\t") > first_line.count(","):
            delim = "\t"
        elif "|" in first_line and first_line.count("|") > first_line.count(","):
            delim = "|"

        reader = csv.reader(io.StringIO(decoded_text), delimiter=delim)
        for r in reader:
            if any(str(c).strip() != "" for c in r):
                raw_table.append(r)

    if not raw_table:
        raise HTTPException(status_code=422, detail="File is empty or contains no readable data rows.")

    # ── 2. Helper Functions for Cell Parsing ─────────────────────────────────
    def parse_cell_date(raw_val) -> Optional[str]:
        if raw_val is None:
            return None
        s_raw = str(raw_val).strip()
        if not s_raw or s_raw.lower() in ("-", "--", "na", "n/a", "nil", "null", "none"):
            return None

        if isinstance(raw_val, (datetime, date)):
            return raw_val.strftime("%Y-%m-%d")

        # Excel serial number (numeric or string)
        try:
            fval = float(s_raw.replace(",", ""))
            if 20000 <= fval <= 65000:
                return (datetime(1899, 12, 30) + timedelta(days=fval)).strftime("%Y-%m-%d")
            # If purely a number < 20000 or > 65000, it's an amount/id, not a date!
            return None
        except (ValueError, OverflowError):
            pass

        # Must contain some date indicator like -, /, ., or letters
        if not any(c in s_raw for c in ("-", "/", ".")) and not any(c.isalpha() for c in s_raw):
            return None

        s = s_raw.split(" ")[0].split("T")[0]
        for fmt in (
            "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d",
            "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y", "%d-%b-%y", "%d/%m/%y", "%d-%m-%y"
        ):
            try:
                return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
            except ValueError:
                pass

        try:
            return date_parser.parse(s_raw, dayfirst=True).strftime("%Y-%m-%d")
        except Exception:
            return None

    def parse_cell_amount(raw_val) -> Optional[float]:
        if raw_val is None:
            return None
        if isinstance(raw_val, (int, float)):
            val = abs(float(raw_val))
            return round(val, 2) if val > 0 else None

        s = str(raw_val).strip()
        if not s or s.lower() in ("-", "--", "na", "n/a", "nil", "null", "none"):
            return None

        # Remove currency symbols and whitespace
        s = re.sub(r'[₹$€£\s]', '', s)
        s = re.sub(r'^(rs\.?|inr)\s*', '', s, flags=re.I)
        s = re.sub(r'\s*(rs\.?|inr)$', '', s, flags=re.I)
        s = re.sub(r'/-\s*$', '', s)
        s = re.sub(r'\s*(dr\.?|debit)$', '', s, flags=re.I)
        s = re.sub(r'\s*(cr\.?|credit)$', '', s, flags=re.I)
        if s.startswith('(') and s.endswith(')'):
            s = s[1:-1]
        s = s.replace(',', '')
        s = s.lstrip('-').rstrip('-')
        try:
            val = float(s)
            return round(abs(val), 2) if val > 0 else None
        except Exception:
            return None

    # ── 3. Identify Header Row & Map Columns ─────────────────────────────────
    def clean_key(val) -> str:
        return re.sub(r'[^a-z0-9]', '', str(val or "").lower())

    def match_alias(header_text: str, target: str) -> bool:
        ckey = clean_key(header_text)
        if not ckey:
            return False
        
        EXACT_MAP = {
            "date": ["date", "txndate", "transactiondate", "valuedate", "vchdate", "postingdate", "transdate", "dt", "tarikh", "datetime", "timestamp", "dateoftransaction", "expensedate"],
            "debit": ["withdrawal", "withdrawalamt", "withdrawalamount", "debit", "debitamt", "debitamount", "dr", "dramt", "dramount"],
            "credit": ["deposit", "depositamt", "depositamount", "credit", "creditamt", "creditamount", "cr", "cramt"],
            "amount": ["amount", "amt", "price", "cost", "total", "spent", "paid", "expense", "expenditure", "rupees", "rs", "inr", "netamount", "transactionamount", "kharcha", "kharche", "rupaye", "totalexpense", "billamount"],
            "description": ["description", "narration", "particulars", "remarks", "details", "desc", "merchant", "payee", "party", "vendor", "name", "title", "item", "purpose", "vivaran", "note", "notes", "transactiondescription", "transactionremarks", "beneficiary"],
            "category": ["category", "cat", "type", "head", "tag", "expensetype", "categoryname"],
            "payment_method": ["paymentmode", "mode", "channel", "paymentmethod", "modeofpayment", "account", "wallet", "source", "bank"]
        }
        targets = EXACT_MAP.get(target, [])
        if ckey in targets:
            return True
        for t in targets:
            if len(t) >= 4 and t in ckey:
                return True
        return False

    col_mapping: dict[str, int] = {}
    best_score = 0
    best_row_idx = -1
    best_map = {}

    scan_limit = min(40, len(raw_table))
    for idx in range(scan_limit):
        row = raw_table[idx]
        temp_map = {}
        for c_idx, cell in enumerate(row):
            if cell is None:
                continue
            text = str(cell).strip()
            if not text:
                continue
            for field in ["date", "debit", "credit", "amount", "description", "category", "payment_method"]:
                if field not in temp_map:
                    if match_alias(text, field):
                        temp_map[field] = c_idx
                        break

        has_amt = "amount" in temp_map or "debit" in temp_map
        has_dt = "date" in temp_map
        has_desc = "description" in temp_map
        score = (2 if has_amt else 0) + (2 if has_dt else 0) + (1 if has_desc else 0) + len(temp_map)

        if score > best_score and has_amt:
            best_score = score
            best_row_idx = idx
            best_map = temp_map

    start_data_idx = 1
    if best_score >= 3:
        col_mapping = best_map
        if "amount" not in col_mapping and "debit" in col_mapping:
            col_mapping["amount"] = col_mapping["debit"]
        start_data_idx = best_row_idx + 1
    else:
        # Fallback: Type inference from cell contents
        max_cols = max(len(r) for r in raw_table) if raw_table else 0
        col_date_scores = [0] * max_cols
        col_amt_scores = [0] * max_cols
        col_text_scores = [0] * max_cols
        first_data_row = -1

        for r_idx, row in enumerate(raw_table):
            for c_idx in range(min(len(row), max_cols)):
                val = row[c_idx]
                if val is None or str(val).strip() == "":
                    continue
                if parse_cell_date(val):
                    col_date_scores[c_idx] += 2
                    if first_data_row == -1: first_data_row = r_idx
                elif parse_cell_amount(val):
                    col_amt_scores[c_idx] += 2
                    if first_data_row == -1: first_data_row = r_idx
                elif isinstance(val, str) and len(val.strip()) > 2:
                    col_text_scores[c_idx] += 1

        inferred_map = {}
        if max_cols > 0:
            best_date = max(range(max_cols), key=lambda i: col_date_scores[i])
            if col_date_scores[best_date] > 0:
                inferred_map["date"] = best_date

            candidates_amt = [i for i in range(max_cols) if i != inferred_map.get("date")]
            if candidates_amt:
                best_amt = max(candidates_amt, key=lambda i: col_amt_scores[i])
                if col_amt_scores[best_amt] > 0:
                    inferred_map["amount"] = best_amt

            candidates_desc = [i for i in range(max_cols) if i not in (inferred_map.get("date"), inferred_map.get("amount"))]
            if candidates_desc:
                best_desc = max(candidates_desc, key=lambda i: col_text_scores[i])
                if col_text_scores[best_desc] > 0:
                    inferred_map["description"] = best_desc

        if "amount" in inferred_map:
            col_mapping = inferred_map
            start_data_idx = max(0, first_data_row)
        else:
            col_mapping = {"date": 0, "amount": 1, "description": 2}
            start_data_idx = 1

    # ── 4. Process Rows ──────────────────────────────────────────────────────
    imported: list[dict] = []
    skipped: list[dict] = []
    affected_months: set[str] = set()

    data_rows = raw_table[start_data_idx:]
    today_str = datetime.utcnow().strftime("%Y-%m-%d")

    credit_col_idx = col_mapping.get("credit")

    for r_idx, row in enumerate(data_rows, start=start_data_idx + 1):
        def get_val(field: str):
            idx = col_mapping.get(field)
            if idx is not None and idx < len(row):
                return row[idx]
            return None

        raw_date = get_val("date")
        raw_amt = get_val("amount")
        raw_desc = get_val("description")
        raw_cat = get_val("category")
        raw_pay = get_val("payment_method")

        amount = parse_cell_amount(raw_amt)

        # Check if this row is a Deposit / Credit row (income, not an expense)
        if amount is None:
            if credit_col_idx is not None and credit_col_idx < len(row):
                credit_amt = parse_cell_amount(row[credit_col_idx])
                if credit_amt is not None and credit_amt > 0:
                    skipped.append({"row": r_idx, "reason": f"Credit / Deposit (income) transaction of ₹{credit_amt:,.2f}"})
                    continue

            skipped.append({"row": r_idx, "reason": f"Invalid or zero amount: '{raw_amt}'"})
            continue

        parsed_date = parse_cell_date(raw_date) or today_str
        desc = str(raw_desc).strip() if raw_desc is not None and str(raw_desc).strip() != "" else "Imported Expense"
        cat = str(raw_cat).strip() if raw_cat is not None and str(raw_cat).strip() != "" else auto_categorize(desc)
        pay_method = str(raw_pay).strip() if raw_pay is not None and str(raw_pay).strip() != "" else "UPI / Bank"

        result = await tool_add_expense(
            db=db,
            user_id=user_id,
            amount=amount,
            category=cat,
            description=desc,
            date_str=parsed_date,
            payment_method=pay_method,
        )
        imported.append(result)
        affected_months.add(parsed_date[:7])

    for month_str in affected_months:
        await audit_and_generate_alerts(db, user_id, month_str)

    # Dispatch mobile phone notification of data fetch
    if imported:
        try:
            from backend.app.services.messaging import notify_data_fetch
            total_imported_val = sum(float(e.get("amount", 0.0)) for e in imported)
            await notify_data_fetch(
                db=db,
                user_id=user_id,
                fetch_source=f"Spreadsheet ({filename})",
                items_count=len(imported),
                total_amount=total_imported_val
            )
        except Exception as e_notify:
            print(f"[Mobile Notification Notice]: {e_notify}")

    if not imported and skipped:
        raise HTTPException(
            status_code=422,
            detail=f"Could not import rows from '{filename}'. {len(skipped)} row(s) had missing or unrecognized data. Reason: {skipped[0]['reason']}. Please ensure columns include Date, Amount (or Debit), and Description."
        )

    return {
        "status": "success",
        "imported": len(imported),
        "skipped": len(skipped),
        "skipped_details": skipped,
        "expenses": imported,
        "message": f"Successfully imported {len(imported)} expense(s) from '{filename}'" + (
            f" ({len(skipped)} non-expense or header row(s) skipped)" if skipped else ""
        )
    }
